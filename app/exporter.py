"""
结果导出器 —— Markdown → Word/Excel/纯文本

前端通过 /api/export 下载，支持三种格式
"""
import io
import re
from typing import List, Dict
from docx import Document
from docx.shared import Pt, Inches

from app.logger import logger


def markdown_to_docx(md_text: str, title: str = "AI Agent 回答") -> bytes:
    """Markdown → Word (.docx)，保留标题/列表/表格格式"""
    try:
        doc = Document()
        doc.styles['Normal'].font.size = Pt(11)

        lines = md_text.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i]

            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif re.match(r'^\d+[.、] ', line):
                doc.add_paragraph(re.sub(r'^\d+[.、] ', '', line), style='List Number')
            elif '|' in line and i + 1 < len(lines) and '---' in lines[i + 1]:
                # 表格
                headers = [c.strip() for c in line.split('|') if c.strip()]
                i += 2  # skip separator
                rows = []
                while i < len(lines) and '|' in lines[i]:
                    rows.append([c.strip() for c in lines[i].split('|') if c.strip()])
                    i += 1
                if headers:
                    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                    table.style = 'Light Grid Accent 1'
                    for j, h in enumerate(headers):
                        table.rows[0].cells[j].text = h
                    for r_idx, row in enumerate(rows):
                        for c_idx, cell in enumerate(row):
                            if c_idx < len(headers):
                                table.rows[r_idx + 1].cells[c_idx].text = cell
                continue
            elif line.strip() and not line.startswith('```'):
                doc.add_paragraph(line)

            i += 1

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()
    except Exception as e:
        logger.error(f"Word导出失败: {e}")
        raise


def markdown_tables_to_excel(md_text: str) -> bytes:
    """提取 Markdown 中的表格 → Excel (.xlsx)"""
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "数据"

        lines = md_text.split('\n')
        i = 0
        row_num = 1
        while i < len(lines):
            line = lines[i]
            if '|' in line and i + 1 < len(lines) and '---' in lines[i + 1]:
                # 表头
                headers = [c.strip() for c in line.split('|') if c.strip()]
                for j, h in enumerate(headers):
                    ws.cell(row=row_num, column=j + 1, value=h)
                row_num += 1
                i += 2
                while i < len(lines) and '|' in lines[i]:
                    cells = [c.strip() for c in lines[i].split('|') if c.strip()]
                    for j, c in enumerate(cells):
                        if j < len(headers):
                            ws.cell(row=row_num, column=j + 1, value=c)
                    row_num += 1
                    i += 1
                continue
            i += 1

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
    except Exception as e:
        logger.error(f"Excel导出失败: {e}")
        raise
