"""
AI Agent 系统 — Streamlit 企业级前端
启动: streamlit run streamlit_app.py --server.port 8501
"""
import streamlit as st
import requests
import json
import time
import base64
from io import BytesIO
from PIL import Image

import os

# ========== 配置 ==========
# Docker Compose 通过 API_BASE 环境变量注入，本地默认 localhost
BASE_URL = os.getenv("API_BASE", "http://localhost:8000").rstrip("/")

st.set_page_config(
    page_title="AI Agent 企业级助手",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ========== 会话状态 ==========
if "token" not in st.session_state:
    st.session_state.token = None
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "session_id" not in st.session_state:
    st.session_state.session_id = f"st-{int(time.time())}"
if "pending_image" not in st.session_state:
    st.session_state.pending_image = None  # base64 图片数据
if "pending_image_preview" not in st.session_state:
    st.session_state.pending_image_preview = None  # 用于显示的预览


@st.cache_data(ttl=3600)
def export_docx(text: str) -> bytes:
    """Markdown → Word（缓存避免闪烁）"""
    try:
        from docx import Document
        import io
        doc = Document()
        for line in text.split('\n'):
            if line.startswith('# '): doc.add_heading(line[2:], level=1)
            elif line.startswith('## '): doc.add_heading(line[3:], level=2)
            elif line.startswith('- '): doc.add_paragraph(line[2:], style='List Bullet')
            elif line.strip(): doc.add_paragraph(line)
        buf = io.BytesIO()
        doc.save(buf); buf.seek(0)
        return buf.read()
    except Exception:
        return text.encode('utf-8')


@st.cache_data(ttl=3600)
def export_xlsx(text: str) -> bytes:
    """Markdown 表格 → Excel（缓存避免闪烁）"""
    try:
        from openpyxl import Workbook
        import io
        wb = Workbook(); ws = wb.active
        row_num = 1
        for line in text.split('\n'):
            if '|' in line and not '---' in line:
                cells = [c.strip() for c in line.split('|') if c.strip()]
                for j, c in enumerate(cells):
                    ws.cell(row=row_num, column=j+1, value=c)
                row_num += 1
        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return buf.read()
    except Exception:
        return text.encode('utf-8')


# ========== 侧边栏 ==========
with st.sidebar:
    st.title("🤖 AI Agent")

    # 认证
    if not st.session_state.token:
        # 先拉租户列表
        tenants = [{"tenant_id": "default", "name": "默认租户(共享)"}]  # fallback
        try:
            resp = requests.get(f"{BASE_URL}/api/tenants", timeout=5)
            if resp.ok:
                tenants = resp.json().get("tenants", tenants)
        except Exception:
            pass

        with st.form("login_form"):
            tenant_options = {t["name"]: t["tenant_id"] for t in tenants}
            selected_name = st.selectbox("选择租户", list(tenant_options.keys()))
            username = st.text_input("用户名")
            password = st.text_input("密码", type="password")
            submitted = st.form_submit_button("登录")
            if submitted:
                try:
                    resp = requests.post(
                        f"{BASE_URL}/api/token",
                        data={
                            "username": username,
                            "password": password,
                            "tenant_id": tenant_options[selected_name],
                        }
                    )
                    if resp.ok:
                        st.session_state.token = resp.json()["access_token"]
                        st.session_state.session_id = f"st-{int(time.time())}"  # 新登录=新会话
                        st.session_state.chat_history = []
                        st.rerun()
                    else:
                        st.error(resp.json().get("detail", "登录失败"))
                except requests.exceptions.RequestException as e:
                    st.warning(f"API 连接失败，请重试")
    else:
        st.success("已登录")
        if st.button("🚪 登出"):
            try:
                requests.post(
                    f"{BASE_URL}/api/revoke",
                    headers={"Authorization": f"Bearer {st.session_state.token}"}
                )
            except Exception:
                pass
            st.session_state.token = None
            st.session_state.chat_history = []; st.session_state.scene = "balanced"
            st.rerun()

    st.divider()

    # 收藏复用——显示问题编辑区
    if st.session_state.get("_reuse_question") and not st.session_state.get("_reuse_shown"):
        st.session_state._reuse_shown = True
        st.info(f"📋 已填入收藏的问题，编辑后发送：")
        reuse_prompt = st.text_input("编辑问题", value=st.session_state._reuse_question, key="reuse_input")
        c1, c2 = st.columns([1, 4])
        with c1:
            if st.button("🚀 发送", type="primary", key="send_reuse"):
                st.session_state._auto_send = reuse_prompt
                st.session_state._reuse_question = None
                st.session_state._reuse_shown = False
                st.rerun()
        with c2:
            if st.button("取消", key="cancel_reuse"):
                st.session_state._reuse_question = None
                st.session_state._reuse_shown = False
                st.rerun()
        st.stop()

    # 功能选择
    mode = st.radio(
        "功能模式",
        ["💬 对话", "📄 知识库查询", "🧠 多Agent协作", "📊 评估"],
    )

    st.divider()

    # 系统状态
    try:
        resp = requests.get(f"{BASE_URL}/api/stats")
        if resp.ok:
            stats = resp.json()
            kb_docs = stats.get("knowledge_base", {}).get("document_count", 0)
            sessions = stats.get("sessions", 0)
            st.metric("📚 文档数", kb_docs)
            st.metric("💬 活跃会话", sessions)
    except requests.exceptions.RequestException:
        st.warning("API 未连接")

    st.divider()

    # 监控
    try:
        resp = requests.get(f"{BASE_URL}/api/metrics", timeout=5)
        if resp.ok:
            with st.expander("📈 监控指标"):
                # 解析关键指标而非原始 dump
                lines = resp.text.split("\n")
                metrics = {}
                for line in lines:
                    if line.startswith("#") or not line.strip():
                        continue
                    # 格式: name{labels} value
                    if "{" in line:
                        name = line.split("{")[0]
                        val = line.rsplit(" ", 1)[-1]
                    else:
                        parts = line.split()
                        name, val = parts[0], parts[-1] if len(parts) > 1 else "?"
                    if name not in metrics:
                        metrics[name] = 0
                    try:
                        metrics[name] += float(val)
                    except ValueError:
                        pass
                if metrics:
                    cols = st.columns(min(len(metrics), 3))
                    for i, (k, v) in enumerate(metrics.items()):
                        with cols[i % 3]:
                            st.metric(k, f"{v:.0f}" if v == int(v) else f"{v:.2f}")
                else:
                    st.caption("暂无指标数据")
    except Exception:
        pass


# ========== 主界面 ==========
st.title("AI Agent 企业级助手")

# ========== 对话模式 ==========
if mode == "💬 对话":
    st.caption("与 AI Agent 自由对话，支持工具调用和知识库检索")

    # 场景选择
    scene_labels = {"precise": "🎯 精准", "balanced": "⚖️ 平衡", "creative": "🎨 创意"}
    if "scene" not in st.session_state:
        st.session_state.scene = "balanced"
    selected_scene_label = st.radio(
        "回答风格", list(scene_labels.values()),
        index=list(scene_labels.keys()).index(st.session_state.scene),
        horizontal=True,
        label_visibility="collapsed"
    )
    st.session_state.scene = {v: k for k, v in scene_labels.items()}[selected_scene_label]

    # 图片上传区域
    with st.expander("🖼️ 上传图片（可选）", expanded=bool(st.session_state.pending_image_preview)):
        uploaded_image = st.file_uploader(
            "选择图片进行分析", type=["jpg", "jpeg", "png", "webp", "gif"],
            label_visibility="collapsed",
            key="image_uploader"
        )
        if uploaded_image:
            # 读取并编码为 base64
            img_bytes = uploaded_image.read()
            img = Image.open(BytesIO(img_bytes))
            # 生成预览
            buf = BytesIO()
            img.thumbnail((300, 300))
            img.save(buf, format="PNG")
            st.image(buf, caption="已选择图片", use_column_width=True)
            # 存 base64
            st.session_state.pending_image = base64.b64encode(img_bytes).decode()
            st.session_state.pending_image_preview = buf.getvalue()
            if st.button("✕ 移除图片"):
                st.session_state.pending_image = None
                st.session_state.pending_image_preview = None
                st.rerun()

    if st.session_state.token:
        for msg in st.session_state.chat_history:
            with st.chat_message(msg["role"]):
                if msg.get("image"):
                    st.image(msg["image"], width=200)
                st.write(msg["content"])
                # 工具按钮（历史消息）
                if msg["role"] == "assistant" and msg["content"]:
                    idx = st.session_state.chat_history.index(msg)
                    c1, c2 = st.columns([0.5, 0.5], gap="small")
                    with c1:
                        st.download_button("📄 Word", data=export_docx(msg["content"]),
                            file_name="answer.docx", key=f"h_docx_{idx}",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    if '|' in msg["content"] and '---' in msg["content"]:
                        with c2:
                            st.download_button("📊 Excel", data=export_xlsx(msg["content"]),
                                file_name="data.xlsx", key=f"h_xlsx_{idx}",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        if prompt := st.chat_input("输入消息..."):
            if not st.session_state.token:
                st.warning("请先在左侧登录后再发送消息")
                st.stop()
            image_data = st.session_state.pending_image
            st.session_state.chat_history.append({
                "role": "user", "content": prompt,
                "image": st.session_state.pending_image_preview
            })
            with st.chat_message("user"):
                if image_data and st.session_state.pending_image_preview:
                    st.image(st.session_state.pending_image_preview, width=200)
                st.write(prompt)
            st.session_state.pending_image = None
            st.session_state.pending_image_preview = None

            with st.chat_message("assistant"):
                placeholder = st.empty()
                full_text = ""
                sources = []
                try:
                    headers = {"Content-Type": "application/json"}
                    if st.session_state.token:
                        headers["Authorization"] = f"Bearer {st.session_state.token}"
                    resp = requests.post(
                        f"{BASE_URL}/api/agent/stream",
                        json={"message": prompt, "session_id": st.session_state.session_id,
                              "mode": "auto", "image": image_data,
                              "scene": st.session_state.get("scene", "balanced")},
                        headers=headers, stream=True, timeout=120)
                    for line in resp.iter_lines():
                        if line and line.startswith(b"data: "):
                            try:
                                event = json.loads(line[6:])
                                if event["type"] == "token":
                                    full_text += event["data"]
                                    placeholder.markdown(full_text + "▌")
                                elif event["type"] == "info":
                                    placeholder.info(event["data"])
                                elif event["type"] == "sources":
                                    sources = event.get("data", [])
                            except json.JSONDecodeError:
                                pass
                    placeholder.markdown(full_text)
                    if sources:
                        with st.expander("📚 参考来源"):
                            for s in sources:
                                st.caption(f"📖 [{s.get('relevance', '')}] {s.get('source', '未知')}")
                    st.session_state.chat_history.append({"role": "assistant", "content": full_text})
                    # 导出按钮
                    cw1, cw2 = st.columns([0.5, 0.5], gap="small")
                    with cw1:
                        st.download_button("📄 Word", data=export_docx(full_text),
                            file_name="answer.docx", key=f"docx_{int(time.time())}",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    if '|' in full_text and '---' in full_text:
                        with cw2:
                            st.download_button("📊 Excel", data=export_xlsx(full_text),
                                file_name="data.xlsx", key=f"xlsx_{int(time.time())}",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                except Exception as e:
                    st.error(f"请求失败: {e}")

        if st.button("🗑️ 清空对话", help="清空当前对话记录（不影响知识库）"):
            st.session_state.chat_history = []
            st.session_state.session_id = f"st-{int(time.time())}"
            st.rerun()
    else:
        st.info("👈 请先在左侧登录后开始对话")

# ========== 知识库查询 ==========
elif mode == "📄 知识库查询":
    st.caption("检索知识库中的文档内容")

    col1, col2 = st.columns([3, 1])
    with col1:
        query = st.text_input("搜索关键词")
    with col2:
        top_k = st.number_input("结果数", 1, 10, 3)

    if query and st.button("搜索"):
        if not st.session_state.token:
            st.warning("请先在左侧登录后使用知识库查询")
        else:
            try:
                headers = {"Authorization": f"Bearer {st.session_state.token}"}
                resp = requests.get(f"{BASE_URL}/api/rag/search", params={"q": query, "top_k": top_k}, headers=headers)
                if resp.ok:
                    results = resp.json()["results"]
                    for i, r in enumerate(results):
                        with st.expander(
                            f"{i+1}. [{r['relevance']}] {r['metadata'].get('source', '未知')} "
                            f"(分数: {r['score']:.2f})"
                        ):
                            st.markdown(r["text"])
                else:
                    st.error(f"搜索失败: {resp.status_code}")
            except Exception as e:
                st.error(f"API 不可达: {e}")

# ========== 多Agent协作 ==========
elif mode == "🧠 多Agent协作":
    st.caption("Supervisor 模式：自动分解复杂任务，多专家协作")

    # 防重复提交 + 结果持久化
    if "ma_running" not in st.session_state:
        st.session_state.ma_running = False
    if "ma_result" not in st.session_state:
        st.session_state.ma_result = None

    with st.form("multi_agent_form"):
        task = st.text_area("描述复杂任务", height=100, placeholder="例如：分析知识库中的卡口系统优势，并提出改进建议")
        col1, col2 = st.columns(2)
        with col1:
            auto = st.checkbox("自动执行", value=False, help="勾选后自动执行，不勾选需人工审批敏感任务")
        submitted = st.form_submit_button("执行任务", type="primary", disabled=st.session_state.ma_running)

    if submitted and task:
        st.session_state.ma_running = True
        with st.spinner("多Agent协作中，预计 20-60 秒..."):
            try:
                resp = requests.post(
                    f"{BASE_URL}/api/multi-agent/run",
                    params={"task": task, "auto_approve": auto},
                    timeout=120,
                )
                if resp.ok:
                    st.session_state.ma_result = resp.json()
                else:
                    st.error(f"执行失败: {resp.text}")
                    st.session_state.ma_result = None
            except Exception as e:
                st.error(f"API 不可达: {e}")
                st.session_state.ma_result = None
        st.session_state.ma_running = False
        st.rerun()

    # 显示上次结果
    if st.session_state.ma_result:
        data = st.session_state.ma_result
        st.subheader("📋 最终回答")
        st.markdown(data["final_answer"])
        st.subheader("🔍 子任务详情")
        for i, sub in enumerate(data.get("subtasks", [])):
            status_icon = {"completed": "✅", "failed": "❌", "running": "🔄"}.get(
                sub["status"], "⏳"
            )
            st.write(f"{status_icon} [{sub['role']}] {sub['description']}")
        if data.get("needs_human_review"):
            st.warning("⚠️ 此任务需要人工审核")

# ========== 评估 ==========
elif mode == "📊 评估":
    st.caption("RAGAS 质量评估：量化检索和生成质量")

    query = st.text_input("测试问题")
    answer = st.text_area("AI 回答", height=150)

    if query and answer and st.button("运行评估"):
        try:
            resp = requests.post(
                f"{BASE_URL}/api/rag/evaluate",
                params={"query": query, "answer": answer},
            )
            if resp.ok:
                data = resp.json()
                col1, col2, col3, col4 = st.columns(4)
                with col1:
                    st.metric("综合分", data.get("overall", 0))
                with col2:
                    st.metric("上下文相关", data.get("context_relevance", {}).get("score", 0))
                with col3:
                    st.metric("忠实度", data.get("faithfulness", {}).get("score", 0))
                with col4:
                    st.metric("回答相关", data.get("answer_relevance", {}).get("score", 0))

                with st.expander("详细报告"):
                    st.json(data)
            else:
                st.error("评估失败")
        except Exception as e:
            st.error(f"API 不可达: {e}")

# ========== 底部 ==========
st.divider()
st.caption(f"会话 ID: {st.session_state.session_id} | API: {BASE_URL}")
